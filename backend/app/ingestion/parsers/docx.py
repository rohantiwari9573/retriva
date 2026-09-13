"""DOCX parsing via python-docx.

One element per paragraph, tagged with the text of the nearest preceding
Heading-style paragraph (best-effort "section" metadata - DOCX has no page
concept at the file-format level, since pagination is computed by the
renderer, not stored in the document)."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.core.config import settings
from app.ingestion.errors import ParsingError
from app.ingestion.parsers.base import ParsedDocument, ParsedElement


class DOCXParser:
    def parse(self, data: bytes) -> ParsedDocument:
        # DOCX is a ZIP archive; python-docx has no built-in cap on how much
        # a small file can decompress to. Sum each entry's declared
        # uncompressed size before python-docx (or anything else) reads a
        # single byte of decompressed content - a classic zip-bomb can be a
        # few KB on disk and gigabytes decompressed.
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                total_uncompressed = sum(info.file_size for info in zf.infolist())
        except zipfile.BadZipFile as exc:
            raise ParsingError(f"Could not read DOCX: {exc}") from exc
        if total_uncompressed > settings.MAX_DOCX_UNCOMPRESSED_SIZE_BYTES:
            raise ParsingError(
                f"DOCX decompresses to {total_uncompressed} bytes, exceeding the "
                f"limit of {settings.MAX_DOCX_UNCOMPRESSED_SIZE_BYTES}."
            )

        try:
            doc = DocxDocument(io.BytesIO(data))
        except (PackageNotFoundError, ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise ParsingError(f"Could not read DOCX: {exc}") from exc

        elements: list[ParsedElement] = []
        current_section: str | None = None

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name if paragraph.style else "") or ""
            if style_name.lower().startswith("heading") or style_name.lower() == "title":
                current_section = text
                continue
            elements.append(ParsedElement(text=text, section=current_section))

        return ParsedDocument(elements=elements)
